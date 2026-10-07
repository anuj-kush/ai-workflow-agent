from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
WORKFLOW_FILE = DATA_DIR / "workflows.xlsx"


def create_csv(filename, rows):
    path = DATA_DIR / filename
    dataframe = pd.DataFrame(rows)
    dataframe.to_csv(path, index=False)
    print(f"Created: {path.name}")


def create_demo_data():
    DATA_DIR.mkdir(exist_ok=True)

    # --------------------------------------------------
    # WF003 - Vendor products
    # --------------------------------------------------

    create_csv(
        "vendor_products.csv",
        [
            {
                "SKU": "V001",
                "Product Name": "Blue Cotton Shirt",
                "Category": "Shirts",
                "Price": 25,
                "Material": "Cotton",
            },
            {
                "SKU": "V002",
                "Product Name": "Black Jeans",
                "Category": "Jeans",
                "Price": 40,
                "Material": "Denim",
            },
            {
                "SKU": "",
                "Product Name": "Running Shoes",
                "Category": "Footwear",
                "Price": 60,
                "Material": "Mesh",
            },
            {
                "SKU": "V004",
                "Product Name": "",
                "Category": "T-Shirts",
                "Price": 18,
                "Material": "Cotton",
            },
            {
                "SKU": "V005",
                "Product Name": "Cotton Hoodie",
                "Category": "Hoodies",
                "Price": 35,
                "Material": "Cotton",
            },
        ],
    )

    # --------------------------------------------------
    # WF005 - Orders
    # --------------------------------------------------

    create_csv(
        "orders.csv",
        [
            {
                "order_id": "ORD-1001",
                "customer_email": "customer1@example.com",
                "customer_name": "Rahul",
                "items": "Blue Cotton Shirt, Black Jeans",
                "status": "Shipped",
            },
            {
                "order_id": "ORD-1002",
                "customer_email": "customer2@example.com",
                "customer_name": "Priya",
                "items": "Running Shoes",
                "status": "Processing",
            },
            {
                "order_id": "ORD003",
                "customer_email": "customer3@example.com",
                "customer_name": "Amit",
                "items": "Cotton Hoodie",
                "status": "Delivered",
            },
        ],
    )

    create_csv(
        "shipments.csv",
        [
            {
                "order_id": "ORD-1001",
                "shipment_status": "In Transit",
                "tracking_number": "TRK10001",
                "carrier": "DemoExpress",
            },
            {
                "order_id": "ORD-1002",
                "shipment_status": "Preparing",
                "tracking_number": "",
                "carrier": "DemoExpress",
            },
            {
                "order_id": "OR-1003",
                "shipment_status": "Delivered",
                "tracking_number": "TRK10003",
                "carrier": "DemoExpress",
            },
        ],
    )

    # --------------------------------------------------
    # WF006 - Product catalog
    # --------------------------------------------------

    create_csv(
        "products.csv",
        [
            {
                "sku": "P001",
                "product_name": "Blue Cotton Shirt",
                "category": "Shirts",
                "material": "Cotton",
                "color": "Blue",
            },
            {
                "sku": "P002",
                "product_name": "Blue Cotton Shirt",
                "category": "Shirts",
                "material": "Cotton",
                "color": "Blue",
            },
            {
                "sku": "P003",
                "product_name": "Black Jeans",
                "category": "Jeans",
                "material": "Denim",
                "color": "Black",
            },
            {
                "sku": "P004",
                "product_name": "Black Denim Jeans",
                "category": "Jeans",
                "material": "Denim",
                "color": "Black",
            },
            {
                "sku": "P005",
                "product_name": "Running Shoes",
                "category": "Footwear",
                "material": "Mesh",
                "color": "White",
            },
        ],
    )

    # --------------------------------------------------
    # WF008 - Keywords
    # --------------------------------------------------

    create_csv(
        "keywords.csv",
        [
            {"keyword": "how to wash cotton shirt"},
            {"keyword": "buy blue cotton shirt"},
            {"keyword": "blue cotton shirt price"},
            {"keyword": "running shoes"},
            {"keyword": "best running shoes for beginners"},
            {"keyword": "black jeans buy online"},
            {"keyword": "company about us"},
            {"keyword": "cotton hoodie"},
            {"keyword": "cotton hoodie price"},
            {"keyword": "how to style jeans"},
        ],
    )

    # --------------------------------------------------
    # WF009 - Employees
    # --------------------------------------------------

    create_csv(
        "employees.csv",
        [
            {
                "employee_id": "E001",
                "employee_name": "Aman",
                "skills": "python,django,rest api",
                "workload": 40,
            },
            {
                "employee_id": "E002",
                "employee_name": "Neha",
                "skills": "python,fastapi,sql",
                "workload": 20,
            },
            {
                "employee_id": "E003",
                "employee_name": "Rohit",
                "skills": "react,javascript,html,css",
                "workload": 35,
            },
            {
                "employee_id": "E004",
                "employee_name": "Sneha",
                "skills": "python,sql,machine learning",
                "workload": 15,
            },
        ],
    )

    # --------------------------------------------------
    # WF010 - Execution logs
    # --------------------------------------------------

    create_csv(
        "execution_logs.csv",
        [
            {
                "workflow_id": "WF001",
                "step_name": "Load inventory",
                "status": "success",
                "execution_time": 1.2,
                "error": "",
            },
            {
                "workflow_id": "WF001",
                "step_name": "Check stock",
                "status": "success",
                "execution_time": 1.5,
                "error": "",
            },
            {
                "workflow_id": "WF002",
                "step_name": "Compare prices",
                "status": "success",
                "execution_time": 2.1,
                "error": "",
            },
            {
                "workflow_id": "WF002",
                "step_name": "Generate report",
                "status": "failed",
                "execution_time": 5.5,
                "error": "Vendor data mismatch",
            },
            {
                "workflow_id": "WF003",
                "step_name": "Validate vendor file",
                "status": "success",
                "execution_time": 1.8,
                "error": "",
            },
            {
                "workflow_id": "WF004",
                "step_name": "Generate description",
                "status": "success",
                "execution_time": 8.5,
                "error": "",
            },
            {
                "workflow_id": "WF005",
                "step_name": "Search order",
                "status": "success",
                "execution_time": 1.1,
                "error": "",
            },
            {
                "workflow_id": "WF006",
                "step_name": "Find duplicates",
                "status": "success",
                "execution_time": 3.4,
                "error": "",
            },
            {
                "workflow_id": "WF007",
                "step_name": "Generate brief",
                "status": "success",
                "execution_time": 7.2,
                "error": "",
            },
            {
                "workflow_id": "WF008",
                "step_name": "Classify keywords",
                "status": "failed",
                "execution_time": 4.2,
                "error": "Invalid keyword row",
            },
        ],
    )


def update_workflows_excel():
    if not WORKFLOW_FILE.exists():
        raise FileNotFoundError(
            f"Workflow file not found: {WORKFLOW_FILE}"
        )

    workbook = load_workbook(WORKFLOW_FILE)
    worksheet = workbook.active

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    required = {
        "Workflow_ID",
        "Steps",
    }

    missing = required - set(headers)

    if missing:
        raise ValueError(
            "Missing Excel columns: "
            + ", ".join(sorted(missing))
        )

    steps = {
        "WF001": (
            "Load inventory → compare current stock with minimum threshold "
            "→ identify low-stock products → calculate reorder quantity "
            "→ generate restock list"
        ),
        "WF002": (
            "Load product prices → load vendor prices → match products by SKU "
            "→ compare internal and vendor prices → calculate percentage "
            "difference → flag exceptions"
        ),
        "WF003": (
            "Read vendor file → detect columns → normalize column names "
            "→ validate required fields → identify invalid rows "
            "→ produce cleaned dataset"
        ),
        "WF004": (
            "Validate required attributes → create product description "
            "→ generate short description → generate SEO title "
            "→ generate meta description"
        ),
        "WF005": (
            "Validate identifier → search order data → retrieve order status "
            "→ retrieve shipment information → summarize current status"
        ),
        "WF006": (
            "Load product catalog → normalize names and SKUs "
            "→ compare identifiers → compare product attributes "
            "→ group likely duplicates → assign confidence"
        ),
        "WF007": (
            "Validate inputs → identify campaign objective "
            "→ summarize products → create messaging "
            "→ create channel recommendations → create campaign checklist"
        ),
        "WF008": (
            "Read keywords → remove duplicates → classify search intent "
            "→ map keywords to categories → identify high-priority keywords "
            "→ export results"
        ),
        "WF009": (
            "Understand task requirements → compare employee skills "
            "→ check current workload → rank candidates "
            "→ select employee → generate assignment summary"
        ),
        "WF010": (
            "Load execution logs → calculate success/failure rate "
            "→ calculate average execution time → identify frequent errors "
            "→ identify slow steps → generate recommendations"
        ),
    }

    workflow_id_column = headers["Workflow_ID"]
    steps_column = headers["Steps"]

    for row in range(2, worksheet.max_row + 1):
        workflow_id = worksheet.cell(
            row=row,
            column=workflow_id_column,
        ).value

        if workflow_id in steps:
            worksheet.cell(
                row=row,
                column=steps_column,
            ).value = steps[workflow_id]

    workbook.save(WORKFLOW_FILE)

    print("Updated workflows.xlsx")


if __name__ == "__main__":
    create_demo_data()
    update_workflows_excel()

    print("\nAssignment demo data setup completed.")